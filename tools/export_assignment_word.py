"""Render the PDF report's shared content as an editable Word document."""

import io
from datetime import datetime, timezone

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Pt, RGBColor
from reportlab.platypus import Image, PageBreak, Paragraph, Preformatted, Spacer, Table


def write_word_report(story, destination, author):
    """Use the same ordered paragraphs, code, tables, and PNGs as the PDF."""
    document = Document()
    section = document.sections[0]
    section.page_width = Pt(612)
    section.page_height = Pt(792)
    section.left_margin = section.right_margin = Pt(48)
    section.top_margin = Pt(46)
    section.bottom_margin = Pt(42)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(9)
    normal.paragraph_format.line_spacing = Pt(14)
    props = document.core_properties
    props.title = "4-bit Quantum Ripple-Carry Adder"
    props.subject = "Module 4: Qiskit code, simulator test cases, and histograms"
    props.author = props.last_modified_by = author
    props.comments = ""
    props.created = props.modified = datetime.now(timezone.utc)
    pending_page_break = False

    def next_paragraph(style=None):
        nonlocal pending_page_break
        paragraph = document.add_paragraph(style=style)
        if pending_page_break:
            paragraph.paragraph_format.page_break_before = True
            pending_page_break = False
        return paragraph

    def paragraph_format(paragraph, style):
        fmt = paragraph.paragraph_format
        fmt.space_before = Pt(getattr(style, "spaceBefore", 0))
        fmt.space_after = Pt(getattr(style, "spaceAfter", 0))
        fmt.line_spacing = Pt(style.leading)
        fmt.left_indent = Pt(getattr(style, "leftIndent", 0))
        fmt.right_indent = Pt(getattr(style, "rightIndent", 0))
        fmt.keep_with_next = bool(getattr(style, "keepWithNext", False))
        fmt.widow_control = False
        paragraph.alignment = {
            0: WD_ALIGN_PARAGRAPH.LEFT,
            1: WD_ALIGN_PARAGRAPH.CENTER,
            2: WD_ALIGN_PARAGRAPH.RIGHT,
            4: WD_ALIGN_PARAGRAPH.JUSTIFY,
        }.get(getattr(style, "alignment", 0), WD_ALIGN_PARAGRAPH.LEFT)

    for item_index, item in enumerate(story):
        if isinstance(item, PageBreak):
            pending_page_break = True
        elif isinstance(item, Paragraph):
            bullet = bool(item.bulletText)
            style_name = (
                "List Bullet" if bullet else
                "Title" if item.style.name == "Title" else
                "Heading 1" if item.style.name == "Heading2" else "Normal"
            )
            paragraph = next_paragraph(style_name)
            paragraph_format(paragraph, item.style)
            if bullet:
                paragraph.paragraph_format.first_line_indent = Pt(-12)
                following = next(
                    (other for other in story[item_index + 1:]
                     if not isinstance(other, Spacer)), None
                )
                paragraph.paragraph_format.keep_with_next = (
                    isinstance(following, Paragraph) and bool(following.bulletText)
                )
            for fragment in item.frags:
                if getattr(fragment, "lineBreak", False):
                    paragraph.add_run().add_break()
                    continue
                run = paragraph.add_run(fragment.text)
                run.font.name = "Arial"
                run.font.size = Pt(fragment.fontSize)
                run.bold = bool(fragment.bold)
                run.italic = bool(fragment.italic)
                if fragment.link:
                    url = fragment.link[0][1]
                    run.font.color.rgb = RGBColor(0, 0, 255)
                    run.underline = True
                    link = OxmlElement("w:hyperlink")
                    relationship = paragraph.part.relate_to(
                        url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True
                    )
                    link.set(qn("r:id"), relationship)
                    link.append(run._r)
                    paragraph._p.append(link)
        elif isinstance(item, Preformatted):
            paragraph = next_paragraph()
            paragraph_format(paragraph, item.style)
            paragraph.paragraph_format.keep_together = False
            paragraph.paragraph_format.space_after = Pt(8)
            run = paragraph.add_run("\n".join(item.lines))
            run.font.name = "Courier New"
            run.font.size = Pt(item.style.fontSize)
        elif isinstance(item, Image):
            paragraph = next_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.space_after = Pt(10)
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.keep_with_next = True
            paragraph.add_run().add_picture(
                io.BytesIO(item.notebook_png),
                width=Pt(item.drawWidth),
                height=Pt(item.drawHeight),
            )
        elif isinstance(item, Table):
            rows = item._cellvalues
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            table.alignment = WD_TABLE_ALIGNMENT.LEFT
            table.style = "Table Grid"
            table.autofit = False
            for column, width in zip(table.columns, item._argW):
                column.width = Pt(width)
            for row_index, row in enumerate(rows):
                for column_index, value in enumerate(row):
                    cell = table.cell(row_index, column_index)
                    cell.width = Pt(item._argW[column_index])
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    paragraph = cell.paragraphs[0]
                    fmt = paragraph.paragraph_format
                    fmt.space_before = fmt.space_after = Pt(5)
                    fmt.line_spacing = Pt(12)
                    run = paragraph.add_run(str(value))
                    run.font.name = "Arial"
                    run.font.size = Pt(9)
                    run.bold = row_index == 0
                    if row_index == 0:
                        shading = OxmlElement("w:shd")
                        shading.set(qn("w:fill"), "E8EEF5")
                        cell._tc.get_or_add_tcPr().append(shading)
            repeat_header = OxmlElement("w:tblHeader")
            table.rows[0]._tr.get_or_add_trPr().append(repeat_header)
            for row in table.rows:
                no_split = OxmlElement("w:cantSplit")
                row._tr.get_or_add_trPr().append(no_split)
        elif isinstance(item, Spacer):
            continue
        else:
            raise TypeError(f"Unsupported report content: {type(item).__name__}")

    document.save(destination)
